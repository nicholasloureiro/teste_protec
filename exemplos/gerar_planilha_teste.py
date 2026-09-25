"""Gera exemplos/cobrancas_teste.xlsx com dados fictícios de cobrança."""

import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

CLIENTES = [
    ("Padaria Pão Dourado Ltda", "12.345.678/0001-90"), ("Mercearia São José ME", "23.456.789/0001-01"),
    ("Auto Peças Irmãos Costa", "34.567.890/0001-12"), ("Clínica Bem Viver", "45.678.901/0001-23"),
    ("Maria Aparecida Souza", "123.456.789-00"), ("João Carlos Lima", "234.567.890-11"),
    ("Farmácia Saúde Total", "56.789.012/0001-34"), ("Construtora Horizonte", "67.890.123/0001-45"),
    ("Ana Paula Ferreira", "345.678.901-22"), ("Restaurante Sabor Mineiro", "78.901.234/0001-56"),
    ("Escola Pequeno Príncipe", "89.012.345/0001-67"), ("Oficina do Zé", "90.123.456/0001-78"),
]
CABECALHO = ["Documento", "Cliente", "CPF/CNPJ", "Emissão", "Vencimento", "Valor (R$)", "Situação"]


def aba(ws, ano, mes, hoje, inicio_doc):
    ws.append(CABECALHO)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3B57")
    for i, (nome, doc) in enumerate(CLIENTES):
        emissao = date(ano, mes, 1) + timedelta(days=random.randint(0, 9))
        vencimento = emissao + timedelta(days=random.choice([10, 15, 20, 30]))
        if vencimento > hoje:
            situacao = "A vencer"
        else:
            situacao = random.choice(["Pago", "Pago", "Vencido"])
        ws.append([f"NF-{inicio_doc + i}", nome, doc, emissao, vencimento,
                   round(random.uniform(150, 8500), 2), situacao])
    for linha in ws.iter_rows(min_row=2):
        linha[3].number_format = linha[4].number_format = "DD/MM/YYYY"
        linha[5].number_format = "#,##0.00"
    for col, largura in zip("ABCDEFG", [12, 30, 22, 12, 12, 14, 12]):
        ws.column_dimensions[col].width = largura


def main():
    random.seed(42)
    hoje = date(2026, 9, 25)
    wb = Workbook()
    wb.active.title = "Agosto 2026"
    aba(wb.active, 2026, 8, hoje, 1001)
    aba(wb.create_sheet("Setembro 2026"), 2026, 9, hoje, 1101)
    saida = Path(__file__).parent / "cobrancas_teste.xlsx"
    wb.save(saida)
    print(saida)


if __name__ == "__main__":
    main()
