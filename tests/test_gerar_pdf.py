import csv
from datetime import date

import pytest
from openpyxl import Workbook
from pypdf import PdfReader
from pytest_bdd import given, parsers, scenarios, then, when

from planilha_pdf import FormatoNaoSuportado, gerar_pdf

scenarios("features/gerar_pdf.feature")

CABECALHO = ["Cliente", "Vencimento", "Valor"]


@pytest.fixture
def contexto(tmp_path):
    return {"dir": tmp_path}


@given(parsers.parse('uma planilha Excel com os clientes "{a}" e "{b}"'))
def planilha_clientes(contexto, a, b):
    wb = Workbook()
    ws = wb.active
    ws.append(CABECALHO)
    ws.append([a, date(2026, 10, 5), 150.0])
    ws.append([b, date(2026, 10, 10), 320.5])
    contexto["entrada"] = contexto["dir"] / "clientes.xlsx"
    wb.save(contexto["entrada"])


@given(parsers.parse('uma planilha Excel com as abas "{a}" e "{b}"'))
def planilha_abas(contexto, a, b):
    wb = Workbook()
    wb.active.title = a
    wb.active.append(CABECALHO)
    wb.active.append(["Cliente X", date(2026, 1, 5), 10.0])
    ws = wb.create_sheet(b)
    ws.append(CABECALHO)
    ws.append(["Cliente Y", date(2026, 2, 5), 20.0])
    contexto["entrada"] = contexto["dir"] / "abas.xlsx"
    wb.save(contexto["entrada"])


@given(parsers.parse('uma planilha CSV com o cliente "{nome}"'))
def planilha_csv(contexto, nome):
    caminho = contexto["dir"] / "clientes.csv"
    with open(caminho, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(CABECALHO)
        w.writerow([nome, "05/10/2026", "99,90"])
    contexto["entrada"] = caminho


@given(parsers.parse('um arquivo "{nome}"'))
def arquivo_qualquer(contexto, nome):
    caminho = contexto["dir"] / nome
    caminho.write_text("conteúdo")
    contexto["entrada"] = caminho


@when("eu gerar o PDF")
def gerar(contexto):
    contexto["saida"] = gerar_pdf(contexto["entrada"])


@when("eu tentar gerar o PDF")
def tentar_gerar(contexto):
    try:
        gerar_pdf(contexto["entrada"])
    except FormatoNaoSuportado as e:
        contexto["erro"] = str(e)


@then("o arquivo PDF é criado")
def pdf_criado(contexto):
    saida = contexto["saida"]
    assert saida.suffix == ".pdf"
    assert saida.read_bytes().startswith(b"%PDF")


@then(parsers.parse('o PDF contém o texto "{texto}"'))
def pdf_contem(contexto, texto):
    reader = PdfReader(contexto["saida"])
    conteudo = "\n".join(p.extract_text() for p in reader.pages)
    assert texto in conteudo


@then(parsers.parse('recebo o erro "{mensagem}"'))
def recebo_erro(contexto, mensagem):
    assert mensagem in contexto.get("erro", "")
