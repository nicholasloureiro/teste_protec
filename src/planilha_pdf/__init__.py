"""Converte planilhas (.xlsx ou .csv) em relatórios PDF."""

import argparse
import csv
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

FORMATOS = {".xlsx", ".xlsm", ".csv"}


class FormatoNaoSuportado(ValueError):
    pass


def ler_planilha(caminho: Path) -> dict[str, list[list]]:
    """Retorna {nome_da_aba: linhas}. CSV vira uma única aba com o nome do arquivo."""
    if caminho.suffix.lower() == ".csv":
        texto = caminho.read_text(encoding="utf-8-sig")
        try:
            dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t")
        except csv.Error:
            dialeto = csv.excel
        return {caminho.stem: list(csv.reader(texto.splitlines(), dialeto))}

    wb = load_workbook(caminho, read_only=True, data_only=True)
    abas = {}
    for ws in wb.worksheets:
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
        abas[ws.title] = [r for r in linhas if any(c not in (None, "") for c in r)]
    wb.close()
    return abas


def formatar(valor) -> str:
    """Formata no padrão brasileiro: datas dd/mm/aaaa e números 1.234,56."""
    if valor is None:
        return ""
    if isinstance(valor, datetime):
        return valor.strftime("%d/%m/%Y %H:%M") if valor.time() != datetime.min.time() else valor.strftime("%d/%m/%Y")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    if isinstance(valor, float):
        return f"{valor:,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return str(valor)


def _tabela(linhas: list[list], largura: float, estilo) -> Table:
    n_colunas = max(len(r) for r in linhas)
    cabecalho = estilo.clone("cabecalho", textColor=colors.white, fontName="Helvetica-Bold")
    dados = [
        [Paragraph(formatar(c), cabecalho if i == 0 else estilo) for c in r] + [""] * (n_colunas - len(r))
        for i, r in enumerate(linhas)
    ]
    tabela = Table(dados, colWidths=[largura / n_colunas] * n_colunas, repeatRows=1)
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1f3b57")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f5f8")]),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#c8d0d8")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    return tabela


def gerar_pdf(entrada: str | Path, saida: str | Path | None = None, titulo: str | None = None) -> Path:
    """Gera um PDF com uma seção (tabela) por aba da planilha. Retorna o caminho do PDF."""
    entrada = Path(entrada)
    if entrada.suffix.lower() not in FORMATOS:
        raise FormatoNaoSuportado(
            f"Formato não suportado: {entrada.suffix or '(sem extensão)'}. Use {', '.join(sorted(FORMATOS))}."
        )
    saida = Path(saida) if saida else entrada.with_suffix(".pdf")

    estilos = getSampleStyleSheet()
    celula = estilos["BodyText"].clone("celula", fontSize=8, leading=10)
    tamanho = landscape(A4)
    margem = 1.5 * cm
    doc = SimpleDocTemplate(str(saida), pagesize=tamanho, leftMargin=margem, rightMargin=margem,
                            topMargin=margem, bottomMargin=margem, title=titulo or entrada.stem)

    def rodape(canvas, doc):
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(margem, margem / 2, f"Gerado em {datetime.now():%d/%m/%Y %H:%M}")
        canvas.drawRightString(tamanho[0] - margem, margem / 2, f"Página {doc.page}")

    elementos = [Paragraph(titulo or entrada.stem, estilos["Title"])]
    abas = [(nome, linhas) for nome, linhas in ler_planilha(entrada).items() if linhas]
    for i, (nome, linhas) in enumerate(abas):
        if i:
            elementos.append(PageBreak())
        elementos += [Paragraph(nome, estilos["Heading2"]), Spacer(1, 0.3 * cm),
                      _tabela(linhas, tamanho[0] - 2 * margem, celula)]
    if not abas:
        elementos.append(Paragraph("A planilha está vazia.", estilos["BodyText"]))

    doc.build(elementos, onFirstPage=rodape, onLaterPages=rodape)
    return saida


def main() -> None:
    parser = argparse.ArgumentParser(description="Converte uma planilha (.xlsx/.csv) em PDF.")
    parser.add_argument("entrada", help="caminho da planilha")
    parser.add_argument("-o", "--saida", help="caminho do PDF (padrão: mesmo nome da planilha)")
    parser.add_argument("-t", "--titulo", help="título do relatório (padrão: nome do arquivo)")
    args = parser.parse_args()
    try:
        print(gerar_pdf(args.entrada, args.saida, args.titulo))
    except (FormatoNaoSuportado, FileNotFoundError) as e:
        sys.exit(f"Erro: {e}")
